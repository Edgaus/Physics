"""Regenera Condiciones_crecimiento_InGaN.docx solo con el borrador del artículo.

No lee la bitácora ni los planes. Esos viven en otros módulos.
"""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor

from indium_model.article import ARTICLE_TITLE, GROWTH_PARAGRAPH, SAMPLES

HERE = Path(__file__).resolve().parent
DEFAULT_OUTPUT = HERE / "Condiciones_crecimiento_InGaN.docx"

HEADERS = (
    "Muestra",
    "Temperatura de los pozos (°C)",
    "Barreras de GaN (°C)",
    "Ciclo Ga-on / Ga-off",
    "Flujo de N₂ (sccm)",
    "Presión MBE (Torr)",
)


def build_document() -> Document:
    document = Document()
    _page(document)
    _title(document, "Condiciones de crecimiento de InGaN cúbico")
    source = document.add_paragraph()
    run = source.add_run("Tabla extraída solo del borrador del artículo.")
    run.italic = True
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    heading = document.add_paragraph()
    label = heading.add_run("Título. ")
    label.bold = True
    label.font.size = Pt(11)
    title = heading.add_run(ARTICLE_TITLE)
    title.italic = True
    title.font.size = Pt(11)

    fragment = document.add_paragraph()
    fragment_label = fragment.add_run("Fragmento de crecimiento")
    fragment_label.bold = True
    fragment_label.font.size = Pt(12)

    quote = document.add_paragraph()
    quote.paragraph_format.left_indent = Cm(0.75)
    quote_run = quote.add_run(GROWTH_PARAGRAPH)
    quote_run.italic = True
    quote_run.font.size = Pt(10)

    table = document.add_table(rows=1, cols=len(HEADERS))
    table.style = "Table Grid"
    table.autofit = True
    for index, header in enumerate(HEADERS):
        cell = table.rows[0].cells[index]
        cell.text = ""
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(header)
        run.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        _shade(cell, "1F4E79")

    for row in SAMPLES:
        values = (
            row.sample,
            row.well_temperature,
            str(row.barrier_GaN_C),
            row.ga_cycle,
            row.n2_sccm,
            str(row.pmbe),
        )
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = ""
            paragraph = cells[index].paragraphs[0]
            run = paragraph.add_run(value)
            run.font.size = Pt(9)
            if index == 0:
                run.bold = True

    return document


def write_docx(path: Path = DEFAULT_OUTPUT) -> Path:
    build_document().save(path)
    return path


def _page(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Mm(297)
    section.page_height = Mm(210)
    section.left_margin = Mm(14)
    section.right_margin = Mm(14)
    section.top_margin = Mm(14)
    section.bottom_margin = Mm(14)


def _title(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(text)
    run.bold = True
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)


def _shade(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    shading.set(qn("w:val"), "clear")
    properties.append(shading)


def main() -> None:
    parser = argparse.ArgumentParser(description="Regenera la tabla Word del borrador.")
    parser.add_argument("-o", "--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    path = write_docx(args.output)
    print(path)


if __name__ == "__main__":
    main()
