"""Las tres fuentes se mantienen separadas y la tabla sale solo del borrador."""

from __future__ import annotations

import ast
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from docx import Document  # noqa: E402

from indium_model.article import GROWTH_PARAGRAPH, SAMPLES  # noqa: E402
from indium_model.bitacora import WELL_SHUTTER_TIMES, by_sample  # noqa: E402
from indium_model.checks import section_5  # noqa: E402
from indium_model.report import informe  # noqa: E402
from indium_model.plans import WELLS_913  # noqa: E402
from indium_model.units import Torr  # noqa: E402
from generate_tabla import build_document, write_docx  # noqa: E402


class RecordTests(unittest.TestCase):
    def test_section_5_holds(self) -> None:
        failed = [text for text, ok in section_5() if not ok]
        self.assertEqual(failed, [])

    def test_quote_is_the_agents_blockquote(self) -> None:
        agents = (ROOT / "agents.md").read_text(encoding="utf-8")
        quoted = [
            line[2:]
            for line in agents.splitlines()
            if line.startswith("> The structures were grown")
        ]
        self.assertEqual(quoted, [GROWTH_PARAGRAPH])

    def test_pressure_keeps_the_written_mantissa(self) -> None:
        self.assertEqual(str(Torr("4.2", -6)), "4.2×10⁻⁶")
        self.assertEqual(str(Torr("2.93", -7)), "2.93×10⁻⁷")
        self.assertEqual(str(Torr("9.7", -6)), "9.7×10⁻⁶")

    def test_report_shows_monolayer_temperatures_once(self) -> None:
        text = informe()
        self.assertIn(
            "Tsus = 840 °C, Tpiro = 721 °C. Duró del segundo 30 al 34 (4 s).",
            text,
        )
        self.assertNotIn("nitrógeno..", text)
        self.assertIn("Sin la pausa final da 5 min 18 s.", text)

    def test_qw1_duration_matches_the_written_time_and_qw2_does_not(self) -> None:
        qw1, qw2, qw3 = WELLS_913
        self.assertEqual(qw1.seconds_omitting_final_pause(), 4 * 60 + 38)
        self.assertNotEqual(qw2.seconds_omitting_final_pause(), 3 * 60 + 54)
        self.assertEqual(qw3.seconds_omitting_final_pause(), 4 * 60 + 38)
        self.assertNotEqual(qw3.written_duration, "4 min 38 s")

    def test_notebook_has_no_well_shutter_times(self) -> None:
        self.assertIsNone(WELL_SHUTTER_TIMES)
        book = by_sample("MME930")
        self.assertFalse(hasattr(book, "in_open_s"))
        self.assertEqual(len(book.crecimiento()), 1)

    def test_table_script_does_not_import_the_other_sources(self) -> None:
        tree = ast.parse((ROOT / "generate_tabla.py").read_text(encoding="utf-8"))
        imported: set[str] = set()
        for node in ast.walk(tree):
            if isinstance(node, ast.Import):
                imported.update(alias.name for alias in node.names)
            elif isinstance(node, ast.ImportFrom) and node.module:
                imported.add(node.module)
        self.assertIn("indium_model.article", imported)
        self.assertNotIn("indium_model.bitacora", imported)
        self.assertNotIn("indium_model.plans", imported)

    def test_docx_contains_the_draft_and_leaves_out_the_notebook(self) -> None:
        document = build_document()
        text = _document_text(document)
        self.assertIn(GROWTH_PARAGRAPH, text)
        self.assertIn("Tabla extraída solo del borrador del artículo.", text)
        for row in SAMPLES:
            self.assertIn(row.sample, text)
            self.assertIn(row.well_temperature, text)
            self.assertIn(row.n2_sccm, text)
            self.assertIn(str(row.pmbe), text)
        self.assertNotIn("9.7", text)
        self.assertNotIn("489", text)
        self.assertNotIn("2.6 %", text)
        self.assertNotIn("bitácora", text.lower())
        self.assertNotIn("Tsus", text)
        self.assertNotIn("170", text)

        with tempfile.TemporaryDirectory() as folder:
            path = write_docx(Path(folder) / "tabla.docx")
            saved = _document_text(Document(path))
        self.assertEqual(saved, text)


def _document_text(document: Document) -> str:
    chunks = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            chunks.extend(cell.text for cell in row.cells)
    return "\n".join(chunks)


if __name__ == "__main__":
    unittest.main()
